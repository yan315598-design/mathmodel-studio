"""docx_presentation_postprocess.py 的行为测试: 现场构造小 docx (手工 oMathPara
段落 + 普通数据表) 走 process() 全链, 断言:

- 编号提出后**公式正文完整** (防"从尾回溯删光公式"回归, 工作区血泪教训);
- 编号跨 run 拆分 (pandoc 把 "(N)" 拆成 '('/')' 多个 m:t) 正确拼接提取;
- H1: 紧贴公式主体的 (整数) 不是编号 (f(1)/a+(2)/g(n) 一字不改), 真编号判据
  = "(" 前存在幸存空白分隔 (\\qquad 渲染产物, 实证 U+2001×2 专属 run);
- H2: 直接格式注入按 OOXML schema 子元素序 (tblPr/trPr/tcPr/pPr/borders),
  对合成与真实 pandoc 表格双重断言;
- 幂等: 第二遍 Step B 0 处理、结构无重复元素、备份不被覆盖;
- Step A **政策门** (v3.1.0): upright 注入 m:sty=p; conventional / 未登记
  保持输入一字不动并警告; 绝不把所有 math run 刷成斜体; 先 upright 后
  conventional 时不撤销不可区分的既有 p 标记而是警示回源重导出;
- Step B 包裹表宽度按**所在节版心** (多节不同页宽各按各节; 非 9072 常量);
- Step C 三线属性 / Step D 自适应+居中+cantSplit / Step E 单元格排版 /
  Step E2 单元格对齐 (短列居中 + 长文本列整列左对齐) 逐属性断言; 公式包裹表
  被 C/D/E/E2 跳过; 相邻公式表间补空段;
- find_unprocessed_equation_numbers 检测口径 (前/后, 与 Step B 同判据)。"""

import contextlib
import io
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

try:
    from docx import Document
    from lxml import etree
    HAS_DEPS = True
except ImportError:
    HAS_DEPS = False

HAS_PANDOC = shutil.which("pandoc") is not None

if HAS_DEPS:
    import docx_presentation_postprocess as _pp
    from docx_presentation_postprocess import (find_unprocessed_equation_numbers,
                                               main, process)
    from docx_presentation_postprocess import (_BORDER_SEQ, _PPR_SEQ, _TCPR_SEQ,
                                               _TBLPR_SEQ, _TRPR_SEQ,
                                               _equation_table_widths)
    # 模块导出的 W/M 是裸命名空间 URI (供 wq()/mq() 包花括号); lxml 查找须用
    # "{URI}tag" 形态, 这里包好供全测试文件使用。
    W = "{%s}" % _pp.W
    M = "{%s}" % _pp.M

# 政策注入 (decision_log dict): 需要 Step A 断言正体的用例统一用它
UPRIGHT = {"math_font": "upright"}
CONVENTIONAL = {"math_font": "conventional"}


def _omp(run_texts):
    """构造 m:oMathPara XML: 每个 run_texts 项一个 m:r/m:t (模拟 pandoc 的
    编号跨 run 拆分, 如 '('、'1'、')')。"""
    runs = "".join(
        f'<m:r><m:t xml:space="preserve">{t}</m:t></m:r>' for t in run_texts)
    return (
        '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
        ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<m:oMath>{runs}</m:oMath></m:oMathPara>"
    )


def _omp_styled(run_specs):
    """构造 m:oMathPara, 每个 run 可带既有 m:sty (模拟 pandoc 对 \\sin/算符/数字
    写 p、字母变量留空的实际产物)。run_specs: [(text, sty | None), ...]。"""
    runs = ""
    for text, sty in run_specs:
        rpr = f'<m:rPr><m:sty m:val="{sty}"/></m:rPr>' if sty else ""
        runs += f'<m:r>{rpr}<m:t xml:space="preserve">{text}</m:t></m:r>'
    return (
        '<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
        ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f"<m:oMath>{runs}</m:oMath></m:oMathPara>"
    )


def _build_fixture(path: Path):
    """01 编号拆 run 的 display 公式; 02 第二条编号公式 (制造相邻公式表);
    03 无编号 display 公式 (Step B 不动、Step A 仍正体); 04 两行数据表。"""
    doc = Document()
    p1 = doc.add_paragraph()
    p1._p.append(etree.fromstring(_omp(["E=mc", "^2", " ", "(", "1", ")"])))
    p2 = doc.add_paragraph()
    p2._p.append(etree.fromstring(_omp(["F=", "ma", "  ", "(", "2", ")"])))
    p3 = doc.add_paragraph()
    p3._p.append(etree.fromstring(_omp(["a", "+", "b"])))
    tbl = doc.add_table(rows=2, cols=2)
    tbl.cell(0, 0).text = "指标"
    tbl.cell(0, 1).text = "数值"
    tbl.cell(1, 0).text = "速度"
    tbl.cell(1, 1).text = "3.0"
    doc.save(str(path))
    return path


def _body(path: Path):
    return Document(str(path)).element.find(W + "body")


def _text_of(el, tag):
    return "".join(t.text or "" for t in el.findall(".//" + tag))


def _eq_tables(body):
    return [t for t in body.findall(W + "tbl")
            if t.find(".//" + M + "oMathPara") is not None]


def _data_tables(body):
    return [t for t in body.findall(W + "tbl")
            if t.find(".//" + M + "oMathPara") is None]


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class ProcessTest(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.docx_path = _build_fixture(Path(self._tmp.name) / "final.docx")

    def tearDown(self):
        self._tmp.cleanup()

    def test_stats_and_backup(self):
        stats = process(self.docx_path, decision_log=UPRIGHT)
        self.assertEqual(1, stats["backup_created"])
        bak = Path(stats["backup_path"])
        self.assertEqual("final.pre_pp.bak.docx", bak.name)
        self.assertTrue(bak.is_file())
        # fixture 共 6+6+3 = 15 个 OMML run 全部正体化
        self.assertEqual(15, stats["upright_runs"])
        self.assertEqual("upright", stats["math_font_policy"])
        self.assertIn("math_font", stats["math_font_source"])
        self.assertEqual(2, stats["equations_wrapped"])
        self.assertEqual(["1", "2"], stats["equation_numbers"])
        self.assertEqual(1, stats["threeline_tables"])
        self.assertEqual(1, stats["fit_tables"])
        # 数据表 4 个单元格各 1 段
        self.assertEqual(4, stats["cell_paragraphs"])
        self.assertEqual(4, stats["cell_aligned"])   # Step E2: 单元格段落全对齐
        self.assertEqual(0, stats["figures_kept"])  # 无图 fixture → Step F 0 处理

    def test_number_extracted_and_formula_body_intact(self):
        """防回归核心: 编号提出后公式正文必须一字不差 (工作区曾把公式删光)。"""
        process(self.docx_path)
        body = _body(self.docx_path)
        eqs = _eq_tables(body)
        self.assertEqual(2, len(eqs))
        for tbl, expect_body, expect_num in zip(eqs, ["E=mc^2", "F=ma"], ["1", "2"]):
            tcs = tbl.findall(W + "tr")[0].findall(W + "tc")
            self.assertEqual(3, len(tcs))
            # 中列: 公式段完整、居中
            self.assertEqual(expect_body, _text_of(tcs[1], M + "t"))
            math_p = tcs[1].find(W + "p")
            jc = math_p.find(W + "pPr/" + W + "jc")
            self.assertEqual("center", jc.get(W + "val"))
            # 右列: 普通文本编号 "(N)"; 段落清零首行缩进 (防编号被挤成两行)
            self.assertEqual(f"({expect_num})", _text_of(tcs[2], W + "t"))
            num_p = tcs[2].find(W + "p")
            ind = num_p.find(W + "pPr/" + W + "ind")
            self.assertEqual("0", ind.get(W + "firstLine"))
            self.assertEqual("0", ind.get(W + "firstLineChars"))
            # 公式正文里不得残留编号痕迹
            self.assertNotIn("(", _text_of(tcs[1], M + "t"))

    def test_all_math_runs_upright(self):
        process(self.docx_path, decision_log=UPRIGHT)
        body = _body(self.docx_path)
        n = 0
        for omath in body.iter(M + "oMath"):
            for r in omath.iter(M + "r"):
                sty = r.find(M + "rPr/" + M + "sty")
                self.assertIsNotNone(sty)
                self.assertEqual("p", sty.get(M + "val"))
                n += 1
        self.assertEqual(15, n)

    def test_equation_tables_keep_borderless_fixed_layout(self):
        """Step C/D/E 必须跳过公式包裹表 (否则三线/autofit 会改坏编号布局)。"""
        process(self.docx_path)
        for tbl in _eq_tables(_body(self.docx_path)):
            borders = tbl.find(W + "tblPr/" + W + "tblBorders")
            for side in ("top", "bottom", "insideH", "insideV"):
                self.assertEqual("none", borders.find(W + side).get(W + "val"), side)
            lay = tbl.find(W + "tblPr/" + W + "tblLayout")
            self.assertEqual("fixed", lay.get(W + "type"))
            self.assertEqual("dxa", tbl.find(W + "tblPr/" + W + "tblW").get(W + "type"))

    def test_blank_paragraph_between_adjacent_equation_tables(self):
        process(self.docx_path)
        children = list(_body(self.docx_path))
        tags = [c.tag for c in children]
        i = tags.index(W + "tbl")  # 第一张公式表
        self.assertEqual(W + "p", tags[i + 1])  # 相邻 w:tbl 间补了空段
        self.assertEqual("", "".join(children[i + 1].itertext()).strip())
        self.assertEqual(W + "tbl", tags[i + 2])

    def test_unnumbered_equation_untouched_by_step_b(self):
        process(self.docx_path, decision_log=UPRIGHT)
        body = _body(self.docx_path)
        body_math_paras = [p for p in body.findall(W + "p")
                           if p.find(".//" + M + "oMathPara") is not None]
        self.assertEqual(1, len(body_math_paras))  # 无编号公式仍是 body 直接子段
        self.assertEqual("a+b", _text_of(body_math_paras[0], M + "t"))
        sty = body_math_paras[0].find(".//" + M + "sty")  # 但 Step A 已正体化
        self.assertEqual("p", sty.get(M + "val"))

    def test_data_table_threeline_attributes(self):
        process(self.docx_path)
        tbl = _data_tables(_body(self.docx_path))[0]
        borders = tbl.find(W + "tblPr/" + W + "tblBorders")
        self.assertEqual("single", borders.find(W + "top").get(W + "val"))
        self.assertEqual("12", borders.find(W + "top").get(W + "sz"))
        self.assertEqual("single", borders.find(W + "bottom").get(W + "val"))
        self.assertEqual("12", borders.find(W + "bottom").get(W + "sz"))
        for side in ("insideH", "insideV", "left", "right"):
            self.assertEqual("none", borders.find(W + side).get(W + "val"), side)
        header_tcs = tbl.findall(W + "tr")[0].findall(W + "tc")
        for tc in header_tcs:
            bottom = tc.find(W + "tcPr/" + W + "tcBorders/" + W + "bottom")
            self.assertEqual("single", bottom.get(W + "val"))
            self.assertEqual("6", bottom.get(W + "sz"))

    def test_data_table_fit_center_and_cantsplit(self):
        process(self.docx_path)
        tbl = _data_tables(_body(self.docx_path))[0]
        tblpr = tbl.find(W + "tblPr")
        tw = tblpr.find(W + "tblW")
        self.assertEqual("5000", tw.get(W + "w"))
        self.assertEqual("pct", tw.get(W + "type"))
        self.assertEqual("center", tblpr.find(W + "jc").get(W + "val"))
        self.assertEqual("autofit", tblpr.find(W + "tblLayout").get(W + "type"))
        for tr in tbl.findall(W + "tr"):
            self.assertIsNotNone(tr.find(W + "trPr/" + W + "cantSplit"))

    def test_data_table_cell_paragraph_spacing(self):
        process(self.docx_path)
        tbl = _data_tables(_body(self.docx_path))[0]
        for p in tbl.findall(".//" + W + "p"):
            sp = p.find(W + "pPr/" + W + "spacing")
            self.assertEqual("0", sp.get(W + "before"))
            self.assertEqual("0", sp.get(W + "after"))
            self.assertEqual("240", sp.get(W + "line"))
            self.assertEqual("auto", sp.get(W + "lineRule"))
            ind = p.find(W + "pPr/" + W + "ind")
            for k in ("firstLine", "firstLineChars", "left", "right"):
                self.assertEqual("0", ind.get(W + k), k)

    def test_rerun_idempotent_and_backup_kept(self):
        process(self.docx_path)
        bak = Path(self.docx_path).with_name("final.pre_pp.bak.docx")
        bak_before = bak.read_bytes()
        stats = process(self.docx_path)
        self.assertEqual(0, stats["equations_wrapped"])  # Step B 第二遍 0 处理
        self.assertEqual([], stats["equation_numbers"])
        self.assertEqual(0, stats["backup_created"])     # 备份存在则不覆盖
        self.assertEqual(bak_before, bak.read_bytes())
        body = _body(self.docx_path)
        self.assertEqual(2, len(_eq_tables(body)))       # 未新增公式表
        # 结构无重复元素 (直接格式覆盖写, 不是追加写)
        for tbl in body.findall(W + "tbl"):
            tblpr = tbl.find(W + "tblPr")
            self.assertEqual(1, len(tblpr.findall(W + "tblBorders")))
            for tr in tbl.findall(W + "tr"):
                trpr = tr.find(W + "trPr")
                if trpr is not None:
                    self.assertEqual(1, len(trpr.findall(W + "cantSplit")))
        # 公式正文在两遍之后仍一字不差
        tc = _eq_tables(body)[0].findall(W + "tr")[0].findall(W + "tc")[1]
        self.assertEqual("E=mc^2", _text_of(tc, M + "t"))

    def test_make_backup_false(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = _build_fixture(Path(tmp) / "x.docx")
            process(p, make_backup=False)
            self.assertEqual([], list(Path(tmp).glob("*.bak.docx")))

    def test_process_missing_file_raises(self):
        with self.assertRaises(FileNotFoundError):
            process(Path(self._tmp.name) / "nope.docx")


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class DetectionTest(unittest.TestCase):
    def test_detects_numbered_display_equations_before_and_after(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = _build_fixture(Path(tmp) / "f.docx")
            self.assertEqual(["1", "2"], find_unprocessed_equation_numbers(p))
            process(p)
            self.assertEqual([], find_unprocessed_equation_numbers(p))

    def test_missing_docx_raises(self):
        with self.assertRaises(Exception):
            find_unprocessed_equation_numbers(Path("___nope___.docx"))


def _assert_schema_ordered(test, parent, seq) -> None:
    """H2: 断言 parent 的子元素标签按 OOXML schema 序严格递增。

    测试 fixture 受控, 只应出现 seq 内的 w: 标签——未知标签按失败处理。
    """
    idx = -1
    for child in parent:
        name = child.tag.split("}")[-1]
        test.assertIn(name, seq, f"未知标签 {name} (parent={parent.tag})")
        pos = seq.index(name)
        test.assertGreater(pos, idx, f"{name} 在 {parent.tag} 内乱序")
        idx = pos


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class UnnumberedFormulaGuardTest(unittest.TestCase):
    """H1 回归: 紧贴公式主体的 "(整数)" 不是编号, 一字不改、不建表。"""

    RUN_SETS = [
        (["f", "(", "1", ")"], "f(1)"),        # pandoc 实测 run 拆分
        (["a", "+", "(", "2", ")"], "a+(2)"),  # $$a+(2)$$ 与 $$a + (2)$$ 同产
        (["g", "(", "n", ")"], "g(n)"),
        (["x", "=", "1", "\u2001\u2001", "(", "1", ")"], None),  # \qquad 真编号
    ]

    def _build(self, path: Path):
        doc = Document()
        for runs, _ in self.RUN_SETS:
            p = doc.add_paragraph()
            p._p.append(etree.fromstring(_omp(runs)))
        doc.save(str(path))
        return path

    def test_only_whitespace_separated_number_extracted(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = self._build(Path(tmp) / "f.docx")
            stats = process(p)
            self.assertEqual(1, stats["equations_wrapped"])
            self.assertEqual(["1"], stats["equation_numbers"])
            body = _body(p)
            paras = [q for q in body.findall(W + "p")
                     if q.find(".//" + M + "oMathPara") is not None]
            self.assertEqual(["f(1)", "a+(2)", "g(n)"],
                             [_text_of(q, M + "t") for q in paras])  # 一字不改
            tcs = _eq_tables(body)[0].findall(W + "tr")[0].findall(W + "tc")
            self.assertEqual("x=1", _text_of(tcs[1], M + "t"))
            self.assertEqual("(1)", _text_of(tcs[2], W + "t"))

    def test_detection_uses_same_criterion(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = self._build(Path(tmp) / "f.docx")
            self.assertEqual(["1"], find_unprocessed_equation_numbers(p))
            process(p)
            self.assertEqual([], find_unprocessed_equation_numbers(p))


def _math_sty_map(path: Path):
    """[(run 文本, m:sty 值 | None), ...] —— 断言"输入是否被改"的语义快照。"""
    out = []
    for omath in _body(path).iter(M + "oMath"):
        for r in omath.iter(M + "r"):
            sty = r.find(M + "rPr/" + M + "sty")
            out.append(("".join(t.text or "" for t in r.findall(M + "t")),
                        sty.get(M + "val") if sty is not None else None))
    return out


def _build_pandoc_like(path: Path, n_eq: int = 6, numbered: bool = True):
    """pandoc 实际产物风格: \\sin/算符/数字带 m:sty=p, 字母变量不带 sty
    (Word 默认渲染为斜体) + 一张三列数据表 (第 3 列长文本)。

    numbered=True 时公式带 "\\qquad (N)" 编号 (供"其余各步照常执行"断言);
    numbered=False 时公式无编号 (Step B 一字不动, 供"输入保持"逐 run 比对)。
    """
    doc = Document()
    for i in range(n_eq):
        runs = [("sin", "p"), (f"x{i}", None), ("+", "p"), ("1", "p")]
        if numbered:
            runs += [("\u2001\u2001", None), ("(", "p"), (str(i + 1), None),
                     (")", "p")]
        p = doc.add_paragraph()
        p._p.append(etree.fromstring(_omp_styled(runs)))
    tbl = doc.add_table(rows=2, cols=3)
    tbl.cell(0, 0).text = "指标"
    tbl.cell(0, 1).text = "数值"
    tbl.cell(0, 2).text = "说明"
    tbl.cell(1, 0).text = "速度"
    tbl.cell(1, 1).text = "3.0"
    tbl.cell(1, 2).text = "跨组对比的补充说明文字超过二十字符阈值说明"
    doc.save(str(path))
    return path


def _set_section_break(p, page_w: int, margin: int = 1440) -> None:
    """把段落 p 变成"节末段"(pPr/sectPr): p 及其之前的内容属本节。

    sectPr 须是 pPr 最后一个子元素 (OOXML schema)。
    """
    ppr = p._p.get_or_add_pPr()
    sect = ppr.makeelement(W + "sectPr", {})
    pgsz = etree.SubElement(sect, W + "pgSz")
    pgsz.set(W + "w", str(page_w))
    pgsz.set(W + "h", "16838")
    mar = etree.SubElement(sect, W + "pgMar")
    for tag in ("top", "right", "bottom", "left"):
        mar.set(W + tag, str(margin))
    mar.set(W + "gutter", "0")
    ppr.append(sect)


def _set_body_section(doc, page_w: int, margin: int = 1440) -> None:
    """改写 body 级 sectPr (文档末节的页宽/边距)。"""
    sect = doc.element.find(W + "body").find(W + "sectPr")
    pgsz = sect.find(W + "pgSz")
    pgsz.set(W + "w", str(page_w))
    pgsz.set(W + "h", "16838")
    mar = sect.find(W + "pgMar")
    for tag in ("top", "right", "bottom", "left"):
        mar.set(W + tag, str(margin))
    mar.set(W + "gutter", "0")


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class MathFontPolicyGateTest(unittest.TestCase):
    """Step A 政策门 (v3.1.0): 尊重 decision_log 政策, 不猜不刷斜体。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.path = _build_pandoc_like(Path(self._tmp.name) / "policy.docx")

    def tearDown(self):
        self._tmp.cleanup()

    def test_unspecified_keeps_input_warns_and_other_steps_still_run(self):
        # 输入保持: 用无编号公式段 (Step B 不动它们) 逐 run 比对文本 + sty
        keep = _build_pandoc_like(Path(self._tmp.name) / "keep.docx",
                                  n_eq=6, numbered=False)
        before = _math_sty_map(keep)
        stats_keep = process(keep, decision_log={})
        self.assertEqual(before, _math_sty_map(keep))  # 一字未改
        self.assertEqual(0, stats_keep["upright_runs"])
        self.assertEqual("unspecified", stats_keep["math_font_policy"])
        self.assertTrue(any("未登记" in w for w in stats_keep["warnings"]))

        # 字体政策缺席不影响其余各步 (步间独立): 带编号 + 数据表的 fixture
        stats = process(self.path, decision_log={})
        self.assertEqual(6, stats["equations_wrapped"])
        self.assertEqual(1, stats["threeline_tables"])
        self.assertEqual(1, stats["fit_tables"])
        self.assertGreater(stats["cell_aligned"], 0)
        self.assertEqual(1, len(stats["long_text_columns"]))

    def test_conventional_keeps_input_and_never_italicizes(self):
        keep = _build_pandoc_like(Path(self._tmp.name) / "keep2.docx",
                                  n_eq=6, numbered=False)
        before = _math_sty_map(keep)
        stats = process(keep, decision_log=CONVENTIONAL)
        after = _math_sty_map(keep)
        self.assertEqual(before, after)                      # 保持输入
        self.assertEqual(0, stats["upright_runs"])
        self.assertEqual("conventional", stats["math_font_policy"])
        # 既有的正体标记 (sin/算符/数字) 必须原样保留
        self.assertEqual("p", dict(after)["sin"])
        # 绝不注入斜体: 全部 sty 只能是 None 或 "p"
        self.assertNotIn("i", {sty for _, sty in after if sty})
        self.assertTrue(any("conventional" in w for w in stats["warnings"]))
        self.assertTrue(any("不等于已实现" in w for w in stats["warnings"]))

    def test_upright_injects_all_runs(self):
        stats = process(self.path, decision_log=UPRIGHT)
        after = _math_sty_map(self.path)
        self.assertEqual(6 * 8, stats["upright_runs"])
        self.assertTrue(all(sty == "p" for _, sty in after))

    def test_upright_then_conventional_keeps_and_warns_reexport(self):
        """先 upright 全刷, 再切 conventional: 不撤销不可区分的覆盖, 而是警示回源。

        (m:sty=p 与 pandoc 自带标记逐 run 不可区分, 逆向擦除会改坏用户在
        \\mathrm/\\text 里写好的正体; 正确动作是 md → docx 回源重导出。)
        """
        process(self.path, decision_log=UPRIGHT)
        after_upright = _math_sty_map(self.path)
        stats2 = process(self.path, decision_log=CONVENTIONAL, make_backup=False)
        self.assertEqual(after_upright, _math_sty_map(self.path))  # 不动文档
        self.assertEqual(0, stats2["upright_runs"])
        flagged, letters = stats2["upright_letter_runs"]
        self.assertGreaterEqual(letters, 5)
        self.assertEqual(letters, flagged)  # 全部字母 run 带 p → 疑似全刷
        self.assertTrue(any("全刷" in w and "回源重导出" in w for w in stats2["warnings"]))

    def test_override_wins_over_decision_log(self):
        stats = process(self.path, decision_log=CONVENTIONAL, math_font="upright")
        self.assertEqual("upright", stats["math_font_policy"])
        self.assertIn("cli", stats["math_font_source"])
        self.assertTrue(all(sty == "p" for _, sty in _math_sty_map(self.path)))

    def test_workspace_decision_log_and_project_root_fallback(self):
        import json
        proj = Path(self._tmp.name) / "proj"
        ws = proj / "paper_workspace"
        ws.mkdir(parents=True, exist_ok=True)
        (proj / "state").mkdir(parents=True, exist_ok=True)
        (proj / "state" / "decision_log.json").write_text(
            json.dumps({"math_font": "upright"}), encoding="utf-8")
        p1 = _build_pandoc_like(ws / "a.docx", n_eq=1)
        stats = process(p1, workspace=ws)              # 骨架约定 <ws>/../state
        self.assertEqual("upright", stats["math_font_policy"])
        p2 = _build_pandoc_like(ws / "b.docx", n_eq=1)
        stats = process(p2, workspace=proj)            # 项目根当 workspace 的兜底
        self.assertEqual("upright", stats["math_font_policy"])
        # 不传 workspace: 按 docx 路径上溯自动定位 (submission/ → proj/state)
        sub = proj / "submission"
        sub.mkdir(parents=True, exist_ok=True)
        p3 = _build_pandoc_like(sub / "c.docx", n_eq=1)
        stats = process(p3)
        self.assertEqual("upright", stats["math_font_policy"])
        self.assertIn("decision_log.json", stats["math_font_source"])  # 定位到的文件

    def test_illegal_decision_log_value_keeps_input(self):
        stats = process(self.path, decision_log={"math_font": "黑体"})
        self.assertEqual(0, stats["upright_runs"])
        self.assertEqual("unspecified", stats["math_font_policy"])
        self.assertTrue(any("无法归一" in w for w in stats["warnings"]))


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class EquationTableWidthTest(unittest.TestCase):
    """Step B 包裹表宽度 = 公式段所在节版心 (v3.1.0 修 9072 写死)。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _num_p(self, doc, num: int):
        p = doc.add_paragraph()
        p._p.append(etree.fromstring(_omp(["x", "=", "1", "\u2001\u2001",
                                           "(", str(num), ")"])))
        return p

    def test_widths_helper_matches_legacy_for_a4(self):
        # A4 11906 − 2×1440 = 9026; 9072 (旧常量) 在同口径下仍得 726/7620/726
        self.assertEqual((726, 7620, 726), _equation_table_widths(9072))
        self.assertEqual(9072, sum(_equation_table_widths(9072)))
        outer, mid, outer2 = _equation_table_widths(10000)
        self.assertEqual((outer, outer2), (726, 726))
        self.assertEqual(10000, outer + mid + outer2)
        # 版心过窄 → 收窄占位列但总宽仍等于版心, 中列 ≥1
        narrow = _equation_table_widths(800)
        self.assertEqual(800, sum(narrow))
        self.assertEqual(240, narrow[0])
        self.assertGreaterEqual(narrow[1], 1)

    def test_multi_section_each_uses_own_text_width(self):
        doc = Document()
        self._num_p(doc, 1)                       # 节 1 (A4: 11906 − 2880 = 9026)
        break_p = doc.add_paragraph("内容分节")
        _set_section_break(break_p, page_w=11906, margin=1440)
        self._num_p(doc, 2)                       # 节 2 (10000 − 2000 = 8000)
        _set_body_section(doc, page_w=10000, margin=1000)
        path = self.dir / "multi.docx"
        doc.save(str(path))

        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual([8000, 9026], stats["equation_widths"])
        self.assertEqual(0, stats["equation_width_fallbacks"])
        eqs = _eq_tables(_body(path))
        self.assertEqual(2, len(eqs))
        for tbl, expect in zip(eqs, (9026, 8000)):
            self.assertEqual("dxa", tbl.find(W + "tblPr/" + W + "tblW").get(W + "type"))
            self.assertEqual(str(expect), tbl.find(W + "tblPr/" + W + "tblW").get(W + "w"))
            grid = [int(g.get(W + "w")) for g in tbl.findall(W + "tblGrid/" + W + "gridCol")]
            self.assertEqual(expect, sum(grid))
            # 公式中列占据剩余宽度 (长公式靠它不顶破版心)
            self.assertEqual(expect - 2 * grid[0], grid[1])

    def test_equation_paragraph_carrying_section_break_uses_own_section(self):
        """节末段本身是公式段: 宽度仍按它结束的那一节算, 但该段跳过不包表
        (节界搬进表格单元格会让 Word 拒绝打开 docx —— Word COM 实测)。"""
        doc = Document()
        self._num_p(doc, 1)                       # 节 1
        p_eq = self._num_p(doc, 2)
        _set_section_break(p_eq, page_w=10000, margin=1000)   # 节 1: 10000−2000
        self._num_p(doc, 3)                       # 末节 (body sectPr)
        _set_body_section(doc, page_w=11000, margin=500)       # 末节: 11000−1000
        path = self.dir / "sec_on_eq.docx"
        doc.save(str(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual(1, stats["equation_section_break_skips"])
        self.assertEqual([8000, 10000], stats["equation_widths"])
        # 公式 (1) 在节 1 → 8000; 公式 (2) 是节末段被跳过; 公式 (3) 在末节 → 10000
        self.assertEqual([8000, 10000], [_tbl_total(t) for t in _eq_tables(_body(path))])
        # 节界仍在 body 直接子层段落上, 没有被藏进任何 tc 内 (文档有效性)
        body = _body(path)
        self.assertIsNotNone(body.find(W + "p/" + W + "pPr/" + W + "sectPr"))
        for tc in body.findall(".//" + W + "tc"):
            self.assertIsNone(tc.find(".//" + W + "pPr/" + W + "sectPr"))
        self.assertTrue(any("节末段" in w for w in stats["warnings"]))

    def test_missing_pgsz_uses_spec_default_page_width(self):
        doc = Document()
        self._num_p(doc, 1)
        sect = doc.element.find(W + "body").find(W + "sectPr")
        pgsz = sect.find(W + "pgSz")
        pgsz.getparent().remove(pgsz)             # 缺 pgSz → OOXML 缺省页宽 12240
        mar = sect.find(W + "pgMar")
        for tag in ("left", "right"):             # 边距显式设 1in → 缺省版心 9360
            mar.set(W + tag, "1440")
        path = self.dir / "no_pgsz.docx"
        doc.save(str(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual([9360], stats["equation_widths"])
        self.assertEqual(0, stats["equation_width_fallbacks"])  # 属性级缺省, 非整体兜底

    def test_no_sectpr_at_all_falls_back_with_warning(self):
        doc = Document()
        self._num_p(doc, 1)
        body = doc.element.find(W + "body")
        body.remove(body.find(W + "sectPr"))      # 连节属性都没有 → 兜底 + 告警
        path = self.dir / "no_sect.docx"
        doc.save(str(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual([9360], stats["equation_widths"])
        self.assertEqual(1, stats["equation_width_fallbacks"])
        self.assertTrue(any("缺省值处理" in w for w in stats["warnings"]))

    def test_multicol_section_gets_unsupported_warning(self):
        """多栏节不建模: 宽度仍按单栏版心算, 但必须发明确"未支持"警告。"""
        doc = Document()
        self._num_p(doc, 1)
        body = doc.element.find(W + "body")
        sect = body.find(W + "sectPr")
        cols = sect.find(W + "cols")            # python-docx 默认模板已带 cols
        if cols is None:                        # (schema 只允许一个 cols, 不能新增)
            cols = etree.SubElement(sect, W + "cols")
        for old in cols.findall(W + "col"):
            cols.remove(old)
        cols.set(W + "num", "2")
        cols.set(W + "space", "425")
        col = etree.SubElement(cols, W + "col")
        col.set(W + "w", "4200")
        col2 = etree.SubElement(cols, W + "col")
        col2.set(W + "w", "4200")
        path = self.dir / "multicol.docx"
        doc.save(str(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual(1, len(stats["equation_width_multicol"]))
        self.assertIn("num=2", stats["equation_width_multicol"][0])
        self.assertTrue(any("多栏版心未支持" in w for w in stats["warnings"]))
        # 仍是单栏版心口径 (12240 − 2×1800 = 8640, python-docx 默认模板边距)
        self.assertEqual([8640], stats["equation_widths"])

    def test_single_column_section_not_warned(self):
        doc = Document()
        self._num_p(doc, 1)
        sect = doc.element.find(W + "body").find(W + "sectPr")
        cols = etree.SubElement(sect, W + "cols")
        cols.set(W + "num", "1")
        path = self.dir / "onecol.docx"
        doc.save(str(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual([], stats["equation_width_multicol"])
        self.assertFalse(any("多栏" in w for w in stats["warnings"]))

    def test_rerun_idempotent_widths(self):
        doc = Document()
        self._num_p(doc, 1)
        _set_body_section(doc, page_w=10000, margin=1000)
        path = self.dir / "rerun.docx"
        doc.save(str(path))
        stats1 = process(path, decision_log=UPRIGHT)
        self.assertEqual([8000], stats1["equation_widths"])
        first = [_tbl_total(t) for t in _eq_tables(_body(path))]
        stats2 = process(path, decision_log=UPRIGHT, make_backup=False)
        self.assertEqual(0, stats2["equations_wrapped"])
        self.assertEqual([], stats2["equation_widths"])   # 第二遍无新包裹表
        self.assertEqual(first, [_tbl_total(t) for t in _eq_tables(_body(path))])
        self.assertEqual(1, len(_eq_tables(_body(path))))  # 未新增包裹表
        self.assertEqual([8000], first)


def _tbl_total(tbl) -> int:
    return sum(int(g.get(W + "w")) for g in tbl.findall(W + "tblGrid/" + W + "gridCol"))


def _cell_aligns(tbl):
    """[[每格首段 jc, ...], ...] (每行一格一项) —— 断言单元格对齐用。"""
    out = []
    for tr in tbl.findall(W + "tr"):
        row = []
        for tc in tr.findall(W + "tc"):
            jcs = []
            for p in tc.findall(W + "p"):
                jc = p.find(W + "pPr/" + W + "jc")
                jcs.append(jc.get(W + "val") if jc is not None else None)
            row.append(jcs[0] if len(jcs) == 1 else jcs)
        out.append(row)
    return out


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class CellAlignTest(unittest.TestCase):
    """Step E2: 单元格段落对齐 (短列居中 / 长文本列整列左对齐) + C/E 步成果保持。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def _build(self, long_text="跨组对比的补充说明文字超过二十字符阈值说明", name="t.docx"):
        doc = Document()
        tbl = doc.add_table(rows=2, cols=3)
        tbl.cell(0, 0).text = "指标"
        tbl.cell(0, 1).text = "数值"
        tbl.cell(0, 2).text = "说明"
        tbl.cell(1, 0).text = "速度"
        tbl.cell(1, 1).text = "3.0"
        tbl.cell(1, 2).text = long_text
        path = self.dir / name
        doc.save(str(path))
        return path

    def test_short_columns_centered_long_column_left(self):
        path = self._build()
        stats = process(path, decision_log=UPRIGHT)
        tbl = _data_tables(_body(path))[0]
        for row in _cell_aligns(tbl):
            self.assertEqual(["center", "center", "left"], row)
        self.assertEqual(6, stats["cell_aligned"])
        self.assertEqual(1, len(stats["long_text_columns"]))
        self.assertIn("第 3 列", stats["long_text_columns"][0])
        self.assertIn("表注说明", stats["long_text_columns"][0])

    def test_no_long_column_means_all_centered(self):
        path = self._build(long_text="短说明", name="short.docx")
        stats = process(path, decision_log=UPRIGHT)
        tbl = _data_tables(_body(path))[0]
        for row in _cell_aligns(tbl):
            self.assertEqual(["center", "center", "center"], row)
        self.assertEqual([], stats["long_text_columns"])

    def test_gridspan_cell_covers_its_columns(self):
        """合并单元格按 gridSpan 覆盖的列区间参与列判定 (两列同时左对齐)。"""
        doc = Document()
        tbl = doc.add_table(rows=3, cols=3)
        tbl.cell(0, 0).text = "指标"
        tbl.cell(0, 1).text = "数值"
        tbl.cell(0, 2).text = "说明"
        tbl.cell(1, 0).text = "速度"
        tbl.cell(1, 1).text = "3.0"
        tbl.cell(1, 2).text = "短"
        row2 = tbl.rows[2]
        merged = row2.cells[0].merge(row2.cells[1])       # 跨第 1-2 列
        merged.text = "跨两列的长文本说明文字超过二十字符阈值说明"
        path = self.dir / "span.docx"
        doc.save(str(path))
        process(path, decision_log=UPRIGHT)
        tbl_el = _data_tables(_body(path))[0]
        row3 = tbl_el.findall(W + "tr")[2].findall(W + "tc")
        jc0 = row3[0].find(W + "p/" + W + "pPr/" + W + "jc")
        self.assertEqual("left", jc0.get(W + "val"))

    def test_threeline_indent_spacing_single_line_preserved(self):
        """E2 只写 jc: C 的三线边框与 E 的零缩进/零段距/单倍行距逐属性保持。"""
        path = self._build()
        process(path, decision_log=UPRIGHT)
        tbl = _data_tables(_body(path))[0]
        borders = tbl.find(W + "tblPr/" + W + "tblBorders")
        self.assertEqual("single", borders.find(W + "top").get(W + "val"))
        self.assertEqual("12", borders.find(W + "top").get(W + "sz"))
        self.assertEqual("single", borders.find(W + "bottom").get(W + "val"))
        for side in ("insideH", "insideV", "left", "right"):
            self.assertEqual("none", borders.find(W + side).get(W + "val"), side)
        for tc in tbl.findall(W + "tr")[0].findall(W + "tc"):
            bottom = tc.find(W + "tcPr/" + W + "tcBorders/" + W + "bottom")
            self.assertEqual("single", bottom.get(W + "val"))
            self.assertEqual("6", bottom.get(W + "sz"))
        for p in tbl.findall(".//" + W + "p"):
            sp = p.find(W + "pPr/" + W + "spacing")
            self.assertEqual(("0", "0", "240", "auto"),
                             (sp.get(W + "before"), sp.get(W + "after"),
                              sp.get(W + "line"), sp.get(W + "lineRule")))
            ind = p.find(W + "pPr/" + W + "ind")
            for k in ("firstLine", "firstLineChars", "left", "leftChars",
                      "right", "rightChars"):
                self.assertEqual("0", ind.get(W + k), k)

    def test_rerun_idempotent_no_duplicate_jc(self):
        path = self._build()
        process(path, decision_log=UPRIGHT)
        before = _cell_aligns(_data_tables(_body(path))[0])
        stats2 = process(path, decision_log=UPRIGHT, make_backup=False)
        self.assertEqual(before, _cell_aligns(_data_tables(_body(path))[0]))
        for p in _data_tables(_body(path))[0].findall(".//" + W + "p"):
            self.assertEqual(1, len(p.findall(W + "pPr/" + W + "jc")))



@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class StandaloneMathTest(unittest.TestCase):
    """S3 前向发现: 独立 m:oMath + 普通文本编号 (非 oMathPara) 不能被 0 处理放行。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    @staticmethod
    def _omp_loose(runs, tail_text):
        """pandoc 少写 display 环境时的形态: m:oMath + 普通文本 "(N)" 同段。"""
        math = "".join(f'<m:r><m:t xml:space="preserve">{t}</m:t></m:r>'
                       for t in runs)
        return ('<m:oMathPara xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
                ' xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>',
                f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math">'
                f"{math}</m:oMath>", tail_text)

    def _build_loose(self, path: Path, tail=" (1)", runs=("T(t)=Ta+(T0-Ta)exp(-kt)",)):
        doc = Document()
        p = doc.add_paragraph("公式说明：T 为温度。")
        p2 = doc.add_paragraph()
        _blank, omath_xml, tail_text = self._omp_loose(list(runs), tail)
        p2._p.append(etree.fromstring(omath_xml))
        r = p2._p.makeelement(W + "r", {})
        t = etree.SubElement(r, W + "t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = tail_text
        p2._p.append(r)
        doc.save(str(path))
        return path

    def test_loose_omath_with_text_number_is_wrapped(self):
        path = self._build_loose(self.dir / "s3.docx")
        # 处理前: 检测器必须认出这个未处理的编号 (否则 docx_to_pdf 不会提醒)
        self.assertEqual(["1"], find_unprocessed_equation_numbers(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual(1, stats["equations_wrapped_standalone"])
        self.assertEqual(0, stats["standalone_math_uncovered"])
        body = _body(path)
        eqs = _eq_tables(body)
        self.assertEqual(1, len(eqs))                     # 已包成编号表
        tcs = eqs[0].findall(W + "tr")[0].findall(W + "tc")
        self.assertEqual("(1)", _text_of(tcs[2], W + "t"))      # 编号进右列
        self.assertEqual("T(t)=Ta+(T0-Ta)exp(-kt)", _text_of(tcs[1], M + "t"))
        # 公式升级为居中 display (不是留成行内), 且原编号文本 run 已删
        para = tcs[1].find(W + "p")
        self.assertIsNotNone(para.find(M + "oMathPara"))
        self.assertEqual("", _text_of(para, W + "t"))
        self.assertEqual([], find_unprocessed_equation_numbers(path))

    def test_wrapper_cells_vertically_centered_and_compact_spacing(self):
        """judge S3 复验 (编号偏低): 三列单元格 vAlign=center, 三列段落同一紧凑行盒。

        任一段落继承样式链的多倍行距/自动间距 → 该列行盒偏高, 垂直居中后编号会
        视觉下移; 故断言三列都显式写了 spacing (0,0,0,0,240,auto)。
        """
        path = self._build_loose(self.dir / "s3_align.docx")
        process(path, decision_log=UPRIGHT)
        for tbl in _eq_tables(_body(path)):
            tcs = tbl.findall(W + "tr")[0].findall(W + "tc")
            self.assertEqual(3, len(tcs))
            for tc in tcs:
                va = tc.find(W + "tcPr/" + W + "vAlign")
                self.assertIsNotNone(va)
                self.assertEqual("center", va.get(W + "val"))
                for p in tc.findall(W + "p"):
                    sp = p.find(W + "pPr/" + W + "spacing")
                    self.assertIsNotNone(sp)
                    self.assertEqual(
                        ("0", "0", "0", "0", "240", "auto"),
                        (sp.get(W + "before"), sp.get(W + "after"),
                         sp.get(W + "beforeLines"), sp.get(W + "afterAutospacing"),
                         sp.get(W + "line"), sp.get(W + "lineRule")))

    def test_loose_omath_tail_number_without_space_is_wrapped(self):
        """review P2: 编号紧贴公式 (x=1(1), 编号前无空白) 也必须走 B2 包表。

        B2 候选判据独立于 Step B 的 H1 空白判据 (那条不改): 只看普通文本尾部的
        "(N)" 形态, 有无空白都不影响。
        """
        path = self._build_loose(self.dir / "s3_nospace.docx", tail="(1)",
                                 runs=("x=1",))
        # 处理前: 检测器也要认出它 (否则 docx_to_pdf 不会提醒)
        self.assertEqual(["1"], find_unprocessed_equation_numbers(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual(1, stats["equations_wrapped_standalone"])
        self.assertEqual(0, stats["standalone_math_uncovered"])
        tcs = _eq_tables(_body(path))[0].findall(W + "tr")[0].findall(W + "tc")
        self.assertEqual("(1)", _text_of(tcs[2], W + "t"))
        self.assertEqual("x=1", _text_of(tcs[1], M + "t"))
        self.assertEqual([], find_unprocessed_equation_numbers(path))

    def test_loose_omath_number_before_formula_kept_and_warned(self):
        """review P2 反例: 编号在公式之前 (无空格) → 原样保留 + 告警, 不搬动。"""
        doc = Document()
        p = doc.add_paragraph()
        r = p._p.makeelement(W + "r", {})
        t = etree.SubElement(r, W + "t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = "(1)"
        p._p.append(r)
        _blank, omath_xml, _tail = self._omp_loose(["x=1"], "")
        p._p.append(etree.fromstring(omath_xml))
        path = self.dir / "before_nospace.docx"
        doc.save(str(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual(0, stats["equations_wrapped_standalone"])
        self.assertEqual(1, stats["standalone_math_uncovered"])
        self.assertEqual(0, len(_eq_tables(_body(path))))          # 未搬动
        body = _body(path)
        self.assertIsNotNone(body.find(W + "p/" + M + "oMath"))     # 公式原位
        self.assertEqual("(1)", "".join(t.text or "" for t in body.iter(W + "t")))

    def test_loose_omath_with_prose_number_warns_not_pretends(self):
        """段内还有说明文字 (普通文本非纯编号) → 只告警, 不假装已呈现合格。"""
        path = self._build_loose(self.dir / "s3b.docx",
                                 tail="见式 (1)", runs=("x=1",))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual(0, stats["equations_wrapped_standalone"])
        self.assertEqual(1, stats["standalone_math_uncovered"])
        self.assertEqual(0, stats["equations_wrapped"])
        self.assertTrue(any("未覆盖" in w and "display 公式" in w
                            for w in stats["warnings"]))
        self.assertEqual(0, len(_eq_tables(_body(path))))   # 未动文档结构

    def test_loose_omath_inline_math_untouched(self):
        """普通行内公式段 (无编号) 一字不动, 也不告警。"""
        doc = Document()
        p = doc.add_paragraph()
        _blank, omath_xml, _tail = self._omp_loose(["x=1"], "")
        p._p.append(etree.fromstring(omath_xml))
        path = self.dir / "inline.docx"
        doc.save(str(path))
        stats = process(path, decision_log=UPRIGHT)
        self.assertEqual(0, stats["equations_wrapped_standalone"])
        self.assertEqual(0, stats["standalone_math_uncovered"])
        self.assertIsNotNone(_body(path).find(W + "p/" + M + "oMath"))

    # ---- 保守结构白名单 (review P2): 复杂形态保持输入 + 告警 ----

    def _build_complex(self, path: Path, kind: str):
        """构造白名单应拒绝的形态; 返回 (path, 期望编号尾部)。"""
        doc = Document()
        p = doc.add_paragraph()
        _blank, omath_xml, _t = self._omp_loose(["x=1"], "")
        omath = etree.fromstring(omath_xml)

        def _text_run(text):
            r = p._p.makeelement(W + "r", {})
            t = etree.SubElement(r, W + "t")
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            t.text = text
            return r

        num = _text_run(" (1)")
        if kind == "hyperlink":
            # 编号在超链里 (旧实现只看直接子 w:r/w:t → 漏判成"无关")
            link = p._p.makeelement(W + "hyperlink", {})
            link.append(num)
            p._p.append(omath)
            p._p.append(link)
        elif kind == "br":
            # 编号 run 里含 w:br (旧实现整条删除 → 连带删掉换行)
            etree.SubElement(num, W + "br")
            p._p.append(omath)
            p._p.append(num)
        elif kind == "before":
            # 编号在公式之前 (旧实现照搬 → 阅读顺序被改)
            p._p.append(num)
            p._p.append(omath)
        elif kind == "two_omath":
            p._p.append(omath)
            p._p.append(etree.fromstring(omath_xml))
            p._p.append(num)
        doc.save(str(path))
        return path

    def test_whitelist_rejects_complex_shapes_keeps_input(self):
        for kind in ("hyperlink", "br", "before", "two_omath"):
            with self.subTest(kind=kind):
                path = self._build_complex(self.dir / f"{kind}.docx", kind)
                texts_before = [t for t, _ in _math_sty_map(path)]
                stats = process(path, decision_log=UPRIGHT)
                self.assertEqual(0, stats["equations_wrapped_standalone"], kind)
                self.assertEqual(1, stats["standalone_math_uncovered"], kind)
                self.assertEqual(0, len(_eq_tables(_body(path))), kind)  # 未搬动
                # 文档语义内容未被删改 (含编号文本仍在, 公式文本一字不差)
                body = _body(path)
                self.assertIsNotNone(body.find(W + "p/" + M + "oMath"), kind)
                text = "".join(t.text or "" for t in body.iter(W + "t"))
                self.assertIn("(1)", text, kind)
                self.assertEqual(
                    texts_before, [t for t, _ in _math_sty_map(path)], kind)


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class IndentSpacingEdgeTest(unittest.TestCase):
    """P2: 悬挂缩进/行数段距/自动间距必须显式清零 (样式继承不得漏过)。"""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_cell_and_equation_paragraphs_clear_hanging_and_autospacing(self):
        doc = Document()
        tbl = doc.add_table(rows=2, cols=2)
        for i, txt in enumerate(("指标", "数值", "速度", "3.0")):
            tbl.cell(i // 2, i % 2).text = txt
        # 直接格式里塞入悬挂缩进 / 行数段距 / 自动间距 (模拟样式继承或前次编辑残留)
        for tc in tbl._tbl.iter(W + "tc"):
            p = tc.find(W + "p")
            ppr = p.find(W + "pPr")
            if ppr is None:
                ppr = p.makeelement(W + "pPr", {})
                p.insert(0, ppr)
            ind = etree.SubElement(ppr, W + "ind")
            ind.set(W + "hanging", "200")
            ind.set(W + "hangingChars", "100")
            ind.set(W + "leftChars", "200")
            sp = etree.SubElement(ppr, W + "spacing")
            sp.set(W + "beforeLines", "50")
            sp.set(W + "afterAutospacing", "1")
        eq = doc.add_paragraph()
        eq._p.append(etree.fromstring(_omp(["x", "=", "1", "\u2001\u2001", "(", "1", ")"])))
        eqppr = eq._p.get_or_add_pPr()
        eqind = etree.SubElement(eqppr, W + "ind")
        eqind.set(W + "hanging", "200")
        eqind.set(W + "hangingChars", "100")
        path = self.dir / "edge.docx"
        doc.save(str(path))

        process(path, decision_log=UPRIGHT)
        body = _body(path)
        for p in body.findall(".//" + W + "p"):
            ind = p.find(W + "pPr/" + W + "ind")
            if ind is None:
                continue
            for k in ("hanging", "hangingChars", "leftChars", "firstLine",
                      "firstLineChars", "left", "right", "rightChars"):
                self.assertEqual("0", ind.get(W + k), (k, ind.attrib))
        tbl_done = _data_tables(body)[0]
        for p in tbl_done.findall(".//" + W + "p"):
            sp = p.find(W + "pPr/" + W + "spacing")
            self.assertEqual(("0", "0", "0", "0", "240", "auto"),
                             (sp.get(W + "before"), sp.get(W + "after"),
                              sp.get(W + "beforeLines"), sp.get(W + "afterAutospacing"),
                              sp.get(W + "line"), sp.get(W + "lineRule")))
        # 公式段 (Step B 中列) 同样清零悬挂缩进
        eq_p = _eq_tables(body)[0].find(W + "tr").findall(W + "tc")[1].find(W + "p")
        self.assertEqual("0", eq_p.find(W + "pPr/" + W + "ind").get(W + "hanging"))



# pandoc 真实产物形态的表格 (tblStyle+tblW+tblLook / trHeight / tcW+vAlign /
# 单元格 pStyle+rPr), 用于 H2 直接格式注入的顺序断言
PANDOC_LIKE_TBL = (
    '<w:tbl xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
    '<w:tblPr><w:tblStyle w:val="Table"/>'
    '<w:tblW w:w="8000" w:type="dxa"/>'
    '<w:tblLook w:val="04A0" w:firstRow="1" w:lastRow="0" '
    'w:firstColumn="1" w:lastColumn="0" w:noHBand="0" w:noVBand="1"/></w:tblPr>'
    '<w:tblGrid><w:gridCol w:w="4000"/><w:gridCol w:w="4000"/></w:tblGrid>'
    '<w:tr><w:trPr><w:trHeight w:val="300"/></w:trPr>'
    '<w:tc><w:tcPr><w:tcW w:w="4000" w:type="dxa"/>'
    '<w:vAlign w:val="center"/></w:tcPr>'
    '<w:p><w:pPr><w:pStyle w:val="Compact"/><w:rPr><w:lang/></w:rPr></w:pPr>'
    '<w:r><w:t>a</w:t></w:r></w:p></w:tc>'
    '<w:tc><w:tcPr><w:tcW w:w="4000" w:type="dxa"/>'
    '<w:vAlign w:val="center"/></w:tcPr>'
    '<w:p><w:pPr><w:pStyle w:val="Compact"/></w:pPr>'
    '<w:r><w:t>b</w:t></w:r></w:p></w:tc>'
    '</w:tr></w:tbl>'
)


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class SchemaOrderTest(unittest.TestCase):
    """H2 回归: C/D/E 的直接格式注入必须按 OOXML schema 子元素序。"""

    def test_pandoc_like_table_schema_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            sect = doc.element.find(W + "body").find(W + "sectPr")
            sect.addprevious(etree.fromstring(PANDOC_LIKE_TBL))
            path = Path(tmp) / "f.docx"
            doc.save(str(path))
            process(path)
            tbl = _data_tables(_body(path))[0]
            _assert_schema_ordered(self, tbl.find(W + "tblPr"), _TBLPR_SEQ)
            _assert_schema_ordered(
                self, tbl.find(W + "tblPr/" + W + "tblBorders"), _BORDER_SEQ)
            for tr in tbl.findall(W + "tr"):
                _assert_schema_ordered(self, tr.find(W + "trPr"), _TRPR_SEQ)
            for tc in tbl.findall(".//" + W + "tc"):
                _assert_schema_ordered(self, tc.find(W + "tcPr"), _TCPR_SEQ)
                ppr = tc.find(W + "p/" + W + "pPr")
                self.assertIsNotNone(ppr)
                _assert_schema_ordered(self, ppr, _PPR_SEQ)
            # 首格 pPr 原有 pStyle+rPr: spacing/ind/jc 必须落在二者之间 (原 bug
            # 两头像过: spacing 插到 pStyle 前 / ind 追加到 rPr 后); jc 是
            # Step E2 的单元格对齐 (短列居中), 按 pPr 序在 ind 之后、rPr 之前
            names = [c.tag.split("}")[-1]
                     for c in tbl.findall(".//" + W + "tc")[0].find(W + "p/" + W + "pPr")]
            self.assertEqual(["pStyle", "spacing", "ind", "jc", "rPr"], names)

    def test_equation_paragraph_ppr_ordered(self):
        """Step B 对既有公式段落注入 ind/jc 也须按 pPr 序 (含 sectPr 尾巴场景)。"""
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            p = doc.add_paragraph()
            p._p.append(etree.fromstring(_omp(["x", "=", "1", "\u2001\u2001",
                                               "(", "3", ")"])))
            path = Path(tmp) / "f.docx"
            doc.save(str(path))
            process(path)
            body = _body(path)
            tc = _eq_tables(body)[0].findall(W + "tr")[0].findall(W + "tc")[1]
            _assert_schema_ordered(self, tc.find(W + "p/" + W + "pPr"), _PPR_SEQ)


@unittest.skipUnless(HAS_DEPS and HAS_PANDOC, "需要 python-docx/lxml 与 pandoc")
class PandocRealityTest(unittest.TestCase):
    """真实 pandoc 产物上的 H1/H2 双重回归 (判据实证来源)。"""

    def _pandoc_docx(self, tmp: str, md_text: str) -> Path:
        md = Path(tmp) / "in.md"
        md.write_text(md_text, encoding="utf-8")
        out = Path(tmp) / "in.docx"
        r = subprocess.run(["pandoc", str(md), "-o", str(out)],
                           capture_output=True, text=True, encoding="utf-8")
        self.assertEqual(0, r.returncode, r.stderr)
        return out

    def test_h1_unnumbered_formulas_untouched(self):
        md = ("$$f(1)$$\n\n$$x=1 \\qquad (1)$$\n\n$$a+(2)$$\n\n"
              "$$a + (2)$$\n\n$$g(n)$$\n")
        with tempfile.TemporaryDirectory() as tmp:
            p = self._pandoc_docx(tmp, md)
            self.assertEqual(["1"], find_unprocessed_equation_numbers(p))
            stats = process(p)
            self.assertEqual(1, stats["equations_wrapped"])
            self.assertEqual(["1"], stats["equation_numbers"])
            body = _body(p)
            texts = sorted(_text_of(q, M + "t") for q in body.findall(W + "p")
                           if q.find(".//" + M + "oMathPara") is not None)
            # 四条无编号公式一字不改 (a + (2) 的空格被 pandoc 丢弃, 也产 a+(2))
            self.assertEqual(["a+(2)", "a+(2)", "f(1)", "g(n)"], texts)
            tcs = _eq_tables(body)[0].findall(W + "tr")[0].findall(W + "tc")
            self.assertEqual("x=1", _text_of(tcs[1], M + "t"))
            self.assertEqual("(1)", _text_of(tcs[2], W + "t"))

    def test_h2_pandoc_table_schema_order(self):
        md = ": 表 1 数据\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
        with tempfile.TemporaryDirectory() as tmp:
            p = self._pandoc_docx(tmp, md)
            process(p)
            for tbl in _body(p).findall(W + "tbl"):
                if tbl.find(".//" + M + "oMathPara") is not None:
                    continue
                tblpr = tbl.find(W + "tblPr")
                _assert_schema_ordered(self, tblpr, _TBLPR_SEQ)
                borders = tblpr.find(W + "tblBorders")
                if borders is not None:
                    _assert_schema_ordered(self, borders, _BORDER_SEQ)
                for tr in tbl.findall(W + "tr"):
                    trpr = tr.find(W + "trPr")
                    if trpr is not None:
                        _assert_schema_ordered(self, trpr, _TRPR_SEQ)
                for tc in tbl.findall(".//" + W + "tc"):
                    _assert_schema_ordered(self, tc.find(W + "tcPr"), _TCPR_SEQ)
                for ppr in tbl.findall(".//" + W + "pPr"):
                    _assert_schema_ordered(self, ppr, _PPR_SEQ)


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class CliTest(unittest.TestCase):
    def test_main_requires_docx_arg(self):
        with self.assertRaises(SystemExit) as ctx:
            with contextlib.redirect_stdout(io.StringIO()):
                with contextlib.redirect_stderr(io.StringIO()):
                    main([])
        self.assertEqual(2, ctx.exception.code)  # argparse required 缺失 → 2

    def test_main_missing_file_exit_2(self):
        buf = io.StringIO()
        with contextlib.redirect_stdout(buf):
            rc = main(["--docx", "___nope___.docx"])
        self.assertEqual(2, rc)
        self.assertIn("不存在", buf.getvalue())

    def test_main_reports_step_stats(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = _build_fixture(Path(tmp) / "f.docx")
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--docx", str(p), "--math-font", "upright"])
            self.assertEqual(0, rc, buf.getvalue())
            out = buf.getvalue()
            self.assertIn("[Step A]", out)
            self.assertIn("全局正体", out)
            self.assertIn("cli --math-font=upright", out)
            self.assertIn("[Step B]", out)
            self.assertIn("编号 1, 2", out)
            self.assertIn("[Step C]", out)
            self.assertIn("[Step E]", out)
            self.assertIn("[Step E2]", out)
            self.assertIn("[Step F]", out)
            self.assertIn("[OK] 已写", out)

    def test_main_without_policy_keeps_input_and_warns(self):
        """不给 --math-font 且就近无 decision_log → 保持输入 + 警告 (exit 仍 0)。"""
        with tempfile.TemporaryDirectory() as tmp:
            p = _build_pandoc_like(Path(tmp) / "f.docx", n_eq=1, numbered=False)
            before = _math_sty_map(p)
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--docx", str(p)])
            self.assertEqual(0, rc, buf.getvalue())
            out = buf.getvalue()
            self.assertIn("未登记", out)
            self.assertIn("保持输入", out)
            self.assertEqual(before, _math_sty_map(p))

    def test_main_math_font_illegal_choice_exit_2(self):
        with self.assertRaises(SystemExit) as ctx:
            with contextlib.redirect_stderr(io.StringIO()):
                main(["--docx", "x.docx", "--math-font", "bold"])
        self.assertEqual(2, ctx.exception.code)  # argparse choices 拦截


@unittest.skipUnless(HAS_DEPS, "需要 python-docx 与 lxml")
class FigureKeepNextTests(unittest.TestCase):
    """E6/Step F: 含图段落 keepNext → 图与其下方图注不被页界劈开。"""

    @staticmethod
    def _png_bytes() -> bytes:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        fig, ax = plt.subplots(figsize=(1, 1))
        ax.plot([0, 1], [0, 1])
        buf = io.BytesIO()
        fig.savefig(buf, format="png")
        plt.close(fig)
        return buf.getvalue()

    def _build(self, path: Path):
        doc = Document()
        doc.add_paragraph("正文段落")
        pic_p = doc.add_paragraph()
        pic_p.add_run().add_picture(io.BytesIO(self._png_bytes()), width=100000)
        doc.add_paragraph("图 1 示意图注")
        doc.add_paragraph("图后正文")
        return doc, path

    def _keep_next_of(self, path: Path):
        body = _body(path)
        return [(p.find(W + "pPr").find(W + "keepNext") is not None
                 if p.find(W + "pPr") is not None else None)
                for p in body.findall(W + "p")]

    def test_image_paragraph_gets_keep_next(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "fig.docx"
            doc, path = self._build(path)
            doc.save(str(path))
            stats = process(path)
            self.assertEqual(1, stats["figures_kept"])
            flags = self._keep_next_of(path)
            # 第 2 段是图片段: keepNext 置位; 正文段与图注段不受影响
            self.assertTrue(flags[1])
            self.assertFalse(any(flags[0:1] + flags[2:]))
            # keepNext 元素按 schema 序落在 pPr 首位附近 (不得排在 rPr 之后)
            ppr = _body(path).findall(W + "p")[1].find(W + "pPr")
            tags = [child.tag.split("}")[1] for child in ppr]
            self.assertLess(tags.index("keepNext"),
                            tags.index("rPr") if "rPr" in tags else len(tags))

    def test_idempotent_second_run_zero(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "fig.docx"
            doc, path = self._build(path)
            doc.save(str(path))
            process(path)
            stats2 = process(path, make_backup=False)
            self.assertEqual(0, stats2["figures_kept"])

    def test_trailing_image_not_counted(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "tail.docx"
            doc = Document()
            doc.add_paragraph("正文")
            pic_p = doc.add_paragraph()
            pic_p.add_run().add_picture(io.BytesIO(self._png_bytes()), width=100000)
            doc.save(str(path))
            stats = process(path)
            self.assertEqual(0, stats["figures_kept"])  # 后无段落, keepNext 无意义


if __name__ == "__main__":
    unittest.main()
