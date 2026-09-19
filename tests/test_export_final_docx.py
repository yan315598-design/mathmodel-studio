"""export_final_docx.py 的行为测试: 编号注入文档序号语义/幂等/冲突、无题注表
警告、围栏 (~~~ 与变长反引号)、标题合成、摘要标题豁免、dry-run 无 pandoc 可跑;
真链路导出 (pandoc + python-docx 齐备时) 验标题块、题注样式与分页落点。"""

import shutil
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
import export_final_docx
from export_final_docx import (NumberingConflict, cleanup_reserved,
                               convert_bibitems, discover_final_files,
                               inject_numbering, main, preprocess_file,
                               reserve_output_path)

HAS_PANDOC = shutil.which("pandoc") is not None
try:
    import docx  # noqa: F401
    HAS_PYTHON_DOCX = True
except ImportError:
    HAS_PYTHON_DOCX = False


def _write(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")


class InjectNumberingTest(unittest.TestCase):
    def test_figures_numbered_in_document_order(self):
        md = ("前文\n\n![甲图](figs/a.png){width=95%}\n\n中间\n\n"
              "![乙图说明](figs/b.png)\n")
        out, stats = inject_numbering(md)
        self.assertEqual(2, stats.n_figures)
        self.assertIn("![图 1 甲图](figs/a.png){width=95%}", out)
        self.assertIn("![图 2 乙图说明](figs/b.png)", out)

    def test_mixed_prefilled_and_new_advance_shared_sequence(self):
        """P1-1 回归: 已编号项占文档序号, 新图不得重号。"""
        md = "![图 1 甲](a.png)\n\n![乙](b.png)\n"
        out, stats = inject_numbering(md)
        self.assertEqual(1, stats.n_figures)
        self.assertEqual(1, stats.n_figures_prefilled)
        self.assertIn("![图 1 甲](a.png)", out)
        self.assertIn("![图 2 乙](b.png)", out)  # 修复前这里错误地产出"图 1"
        self.assertEqual(3, stats.fig_seq)

    def test_full_rerun_is_idempotent(self):
        """对已按文档序号编完号的 md 重跑, 原文不变、只计 prefilled。"""
        md = "![图 1 甲](a.png)\n\n![图 2 乙](b.png)\n"
        out, stats = inject_numbering(md)
        self.assertEqual(0, stats.n_figures)
        self.assertEqual(2, stats.n_figures_prefilled)
        self.assertEqual(md, out)

    def test_prefilled_number_mismatch_raises(self):
        """P1-1 回归: 已有编号与文档序号不一致时报错, 不静默重号。"""
        with self.assertRaises(NumberingConflict):
            inject_numbering("![图 3 甲](a.png)\n")

    def test_empty_alt_skipped_without_consuming_number(self):
        md = "![](b.png)\n\n![新图](c.png)\n"
        out, stats = inject_numbering(md)
        self.assertEqual(1, stats.n_figures)
        self.assertIn("![](b.png)", out)
        self.assertIn("![图 1 新图](c.png)", out)

    def test_table_captions_numbered_with_blank_line_gap(self):
        # 题注与表之间隔空行是本仓库 md 约定 (pandoc 挂接口径)
        md = ": 温度表\n\n| a | b |\n|---|---|\n| 1 | 2 |\n\n: 水分表\n\n| c | d |\n|---|---|\n| 3 | 4 |\n"
        out, stats = inject_numbering(md)
        self.assertEqual(2, stats.n_tables)
        self.assertIn(": 表 1 温度表", out)
        self.assertIn(": 表 2 水分表", out)

    def test_mixed_table_caption_prefilled_advances_sequence(self):
        """P1-1 回归 (表侧): 已有 "表 1" 占号, 新表题注得 "表 2"。"""
        md = ": 表 1 已有\n\n| a |\n|---|\n| 1 |\n\n: 新表\n\n| b |\n|---|\n| 2 |\n"
        out, stats = inject_numbering(md)
        self.assertEqual(1, stats.n_tables)
        self.assertEqual(1, stats.n_tables_prefilled)
        self.assertIn(": 表 1 已有", out)
        self.assertIn(": 表 2 新表", out)

    def test_table_caption_full_rerun_idempotent(self):
        md = ": 表 1 甲\n\n| a |\n|---|\n| 1 |\n\n: 表 2 乙\n\n| b |\n|---|\n| 2 |\n"
        out, stats = inject_numbering(md)
        self.assertEqual(0, stats.n_tables)
        self.assertEqual(2, stats.n_tables_prefilled)
        self.assertEqual(md, out)

    def test_table_caption_number_mismatch_raises(self):
        with self.assertRaises(NumberingConflict):
            inject_numbering(": 表 2 突进\n\n| a |\n|---|\n| 1 |\n")

    def test_uncaptioned_table_collects_warning(self):
        md = "正文段落。\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
        out, stats = inject_numbering(md)
        self.assertEqual(0, stats.n_tables)
        self.assertEqual(1, len(stats.uncaptioned_tables))
        self.assertIn("| a | b |", stats.uncaptioned_tables[0][1])
        self.assertNotIn("表 1", out)

    def test_backtick_fence_content_not_numbered(self):
        md = "```python\n![ fenced ](x.png)\n```\n\n![真实图](y.png)\n"
        out, stats = inject_numbering(md)
        self.assertEqual(1, stats.n_figures)
        self.assertIn("![ fenced ](x.png)", out)
        self.assertIn("![图 1 真实图](y.png)", out)

    def test_tilde_fence_content_not_numbered(self):
        """P2-7 回归: ~~~ 围栏内的图不注入。"""
        md = "~~~\n![ fenced ](x.png)\n~~~\n\n![真实图](y.png)\n"
        out, stats = inject_numbering(md)
        self.assertEqual(1, stats.n_figures)
        self.assertIn("![ fenced ](x.png)", out)
        self.assertIn("![图 1 真实图](y.png)", out)

    def test_longer_backtick_fence_not_closed_by_shorter(self):
        """P2-7 回归: ```` 开栏时其中的 ``` 是内容, 不提前闭栏。"""
        md = "````\n![ fenced ](x.png)\n```\nstill fenced ![ b ](z.png)\n````\n\n![真实图](y.png)\n"
        out, stats = inject_numbering(md)
        self.assertEqual(1, stats.n_figures)
        self.assertIn("still fenced ![ b ](z.png)", out)


class PreprocessFileTest(unittest.TestCase):
    def test_bibitems_converted_with_running_numbers(self):
        text, n, warns = convert_bibitems(
            "\\bibitem{r1} A 论文。\n\n\\bibitem{r2} B 论文。\n")
        self.assertEqual(2, n)
        self.assertEqual([], warns)
        self.assertIn("[1] A 论文。", text)
        self.assertIn("[2] B 论文。", text)

    def test_mixed_explicit_and_bibitem_numbering_matches_audit(self):
        """P1 回归: 显式 [1] Alpha + \\bibitem{b} Beta 混排不得重号, 且与审计一致。"""
        mixed = "[1] Alpha 条目。\n\n\\bibitem{b} Beta 条目。\n"
        text, n, warns = convert_bibitems(mixed)
        self.assertEqual(1, n)
        self.assertEqual([], warns)
        self.assertIn("[1] Alpha 条目。", text)     # 显式条目一字不动
        self.assertIn("[2] Beta 条目。", text)      # bibitem 取未被占用的 2
        self.assertNotIn("[1] Beta", text)          # 修复前这里正是重号
        # 审计端用同一分配函数 → defined 集合与导出后的编号一致
        import ref_order_audit
        with tempfile.TemporaryDirectory() as tmp:
            refs = Path(tmp) / "09_references.md"
            body = Path(tmp) / "02_body.md"
            _write(refs, "## 参考文献\n\n" + mixed)
            _write(body, "## 模型\n\n文献[1]与[2]。\n")
            report = ref_order_audit.audit_ref_order(
                [("02_body.md", body), ("09_references.md", refs)])
        self.assertEqual([1, 2], report["defined_numbers"])
        self.assertEqual(0, report["n_fail"])

    def test_duplicate_explicit_number_warns(self):
        _text, n, warns = convert_bibitems(
            "[1] Alpha。\n\n\\bibitem{b} Beta。\n\n[1] Gamma。\n")
        self.assertEqual(1, n)
        self.assertEqual(1, len(warns))
        self.assertIn("重复", warns[0])

    def test_body_file_embedded_refs_section_converted_in_section_only(self):
        """main.md 单文件稿: 只转 "## 参考文献" 节内的 bibitem, 节外一字不动。"""
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "main.md"
            _write(p, "\n".join([
                "# 论文",
                "正文如文献[1]与[2]所示。",
                "",
                "## 参考文献 {-}",
                "",
                "[1] Alpha。",
                "",
                "\\bibitem{b} Beta。",
                "",
                "## 附录 A",
                "",
                "\\bibitem{stray} 附录里的伪条目",
            ]))
            out, notes = preprocess_file(p, p.read_text(encoding="utf-8"))
        self.assertIn("[1] Alpha。", out)
        self.assertIn("[2] Beta。", out)                    # 节内转换 + 共享分配
        self.assertNotIn("[3] 附录里的伪条目", out)
        self.assertIn("\\bibitem{stray} 附录里的伪条目", out)  # 节外不转
        self.assertTrue(any("正文内嵌参考文献节" in n for n in notes), notes)
        self.assertTrue(any("不在任何" in n and "未转换" in n for n in notes), notes)

    def test_convert_in_refs_sections_warns_without_section(self):
        """无 "## 参考文献" 节但出现 bibitem → 明确告警未转换 (不静默留原命令)。"""
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "main.md"
            _write(p, "## 模型\n\n\\bibitem{x} 没有节头的条目\n")
            _out, notes = preprocess_file(p, p.read_text(encoding="utf-8"))
        self.assertTrue(any("未转换" in n and "## 参考文献" in n for n in notes), notes)

    def test_references_file_gets_synthesized_unnumbered_heading(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "09_references.md"
            _write(p, "\\bibitem{r1} A。\n")
            out, notes = preprocess_file(p, p.read_text(encoding="utf-8"))
            self.assertIn("## 参考文献 {-}", out)
            self.assertTrue(any("参考文献" in n for n in notes))

    def test_appendix_file_headings_all_unnumbered(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "10_appendix.md"
            _write(p, "## 附录A 代码\n\n正文\n")
            out, _ = preprocess_file(p, p.read_text(encoding="utf-8"))
            self.assertIn("## 附录A 代码 {-}", out)

    def test_body_file_special_headings_unnumbered(self):
        with tempfile.TemporaryDirectory() as tmp:
            p = Path(tmp) / "06_models.md"
            _write(p, "## 模型建立\n\n正文\n\n## 附录说明\n\n## 摘 要\n")
            out, _ = preprocess_file(p, p.read_text(encoding="utf-8"))
            self.assertIn("## 模型建立\n", out)  # 正文标题不加标记
            self.assertIn("## 附录说明 {-}", out)
            self.assertIn("## 摘 要 {-}", out)  # P2-8: md 自带摘要标题也豁免


class DiscoverTest(unittest.TestCase):
    def test_numeric_series_natural_order(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp)
            for name in ("10_appendix.md", "01_abstract.md", "02_body.md"):
                _write(ws / name, "x\n")
            found = discover_final_files(ws)
            self.assertEqual(["01_abstract.md", "02_body.md", "10_appendix.md"],
                             [p.name for p in found])

    def test_non_export_md_excluded(self):
        """v3.1.0: 非导出 md (notes/README/草稿) 不进拼接清单 (与引用审计同源)。"""
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp)
            for name in ("10_appendix.md", "01_abstract.md", "notes.md",
                         "README.md", "draft_03_ideas.md"):
                _write(ws / name, "x\n")
            found = discover_final_files(ws)
            self.assertEqual(["01_abstract.md", "10_appendix.md"],
                             [p.name for p in found])


class ReserveOutputPathTest(unittest.TestCase):
    """v3.1.0: 秒级时间戳 + O_EXCL 独占占位 (同秒/并发导出不覆盖)。"""

    def test_second_call_gets_suffix(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            p1 = reserve_output_path(d, "cumcm_final_20260919_120000")
            p2 = reserve_output_path(d, "cumcm_final_20260919_120000")
            p3 = reserve_output_path(d, "cumcm_final_20260919_120000")
            self.assertEqual("cumcm_final_20260919_120000.docx", p1.name)
            self.assertEqual("cumcm_final_20260919_120000_2.docx", p2.name)
            self.assertEqual("cumcm_final_20260919_120000_3.docx", p3.name)
            self.assertEqual([0, 0, 0], [p.stat().st_size for p in (p1, p2, p3)])
            self.assertEqual(["cumcm_final_20260919_120000.docx",
                              "cumcm_final_20260919_120000_2.docx",
                              "cumcm_final_20260919_120000_3.docx"],
                             sorted(p.name for p in d.iterdir()))

    def test_existing_file_never_overwritten(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            existing = d / "x_final_20260919_120000.docx"
            existing.write_bytes(b"OLD")
            got = reserve_output_path(d, "x_final_20260919_120000")
            self.assertEqual("x_final_20260919_120000_2.docx", got.name)
            self.assertEqual(b"OLD", existing.read_bytes())

    def test_cleanup_only_removes_empty_reservation(self):
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp)
            empty = reserve_output_path(d, "a_final_1")
            cleanup_reserved(empty)
            self.assertFalse(empty.exists())
            filled = reserve_output_path(d, "b_final_1")
            filled.write_bytes(b"content")          # 已有内容 → 保留 (pandoc 产物)
            cleanup_reserved(filled)
            self.assertTrue(filled.is_file())

    def test_timestamp_is_second_resolution(self):
        """dry-run 打印的目标名是秒级时间戳 (YYYYMMDD_HHMMSS)。"""
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp) / "paper_workspace"
            _write(ws / "01_abstract.md", "摘要。\n")
            import contextlib, io, re
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--workspace", str(ws), "--out-dir",
                           str(Path(tmp) / "submission"), "--dry-run"])
            self.assertEqual(0, rc)
            self.assertRegex(buf.getvalue(), r"_final_\d{8}_\d{6}\.docx")


class MainDryRunTest(unittest.TestCase):
    """dry-run 不产 docx、不需要 pandoc/python-docx, 返回 0。"""

    def test_dry_run_prints_stats_and_command(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp) / "paper_workspace"
            _write(ws / "01_abstract.md", "摘要正文。\n\n**关键词**：a；b\n")
            _write(ws / "02_body.md", "## 一章\n\n![图说明](x.png)\n\n: 表题\n\n| a |\n|---|\n| 1 |\n")
            out_dir = Path(tmp) / "submission"
            import contextlib, io
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--workspace", str(ws), "--out-dir", str(out_dir), "--dry-run"])
            self.assertEqual(0, rc)
            text = buf.getvalue()
            self.assertIn("将拼接 2 个 md", text)
            self.assertIn("图 1 张", text)
            self.assertIn("表 1 张", text)
            self.assertIn("pandoc", text)
            self.assertFalse(out_dir.exists())  # dry-run 不建目录不写文件

    def test_numbering_conflict_fails_dry_run(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp) / "paper_workspace"
            _write(ws / "01_abstract.md", "![图 3 突进](x.png)\n")
            import contextlib, io
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--workspace", str(ws), "--dry-run"])
            self.assertEqual(2, rc)
            self.assertIn("编号冲突", buf.getvalue())


@unittest.skipUnless(HAS_PANDOC and HAS_PYTHON_DOCX,
                     "需要 pandoc 与 python-docx")
class MainExportTest(unittest.TestCase):
    def test_export_produces_docx_with_title_block_and_captions(self):
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp) / "paper_workspace"
            out_dir = Path(tmp) / "submission"
            _write(ws / "01_abstract.md", "摘要正文一段。\n\n**关键词**：甲；乙\n")
            _write(ws / "02_body.md",
                   "## 问题重述\n\n正文。\n\n![路线图](missing.png)\n\n"
                   ": 结果表\n\n| a | b |\n|---|---|\n| 1 | 2 |\n")
            import contextlib, io
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--workspace", str(ws), "--out-dir", str(out_dir),
                           "--title", "测试论文标题", "--problem-no", "A"])
            self.assertEqual(0, rc, buf.getvalue())
            produced = list(out_dir.glob("*_final_*.docx"))
            self.assertEqual(1, len(produced))
            doc = docx.Document(str(produced[0]))
            texts = [p.text for p in doc.paragraphs]
            self.assertEqual("测试论文标题", texts[0])
            self.assertIn("题号：A", texts[1])
            self.assertIn("摘　要", texts)
            first_h1 = next(p for p in doc.paragraphs if p.style.name == "Heading 1")
            self.assertTrue(first_h1.paragraph_format.page_break_before)
            captions = [p.text for p in doc.paragraphs
                        if p.style.name in ("Image Caption", "Table Caption")]
            self.assertTrue(any(t.startswith("图 1 ") for t in captions), captions)
            self.assertTrue(any(t.startswith("表 1 ") for t in captions), captions)

    def test_own_abstract_heading_unnumbered_and_break_on_body(self):
        """P2-8 回归: md 自带 '## 摘要' 时——摘要标题不编号、不重复插 '摘　要' 段,
        分页落在第一个正文章标题而非摘要标题前。"""
        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp) / "paper_workspace"
            out_dir = Path(tmp) / "submission"
            _write(ws / "01_abstract.md",
                   "## 摘要\n\n摘要正文。\n\n**关键词**：甲；乙\n")
            _write(ws / "02_body.md", "## 问题重述\n\n正文。\n")
            import contextlib, io
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--workspace", str(ws), "--out-dir", str(out_dir),
                           "--title", "T"])
            self.assertEqual(0, rc, buf.getvalue())
            doc = docx.Document(str(list(out_dir.glob("*_final_*.docx"))[0]))
            texts = [p.text for p in doc.paragraphs]
            self.assertNotIn("摘　要", texts)  # 已有摘要标题, 不重复插
            h1s = [p for p in doc.paragraphs if p.style.name == "Heading 1"]
            self.assertEqual(2, len(h1s))
            self.assertEqual("摘要", h1s[0].text)  # 未被 pandoc 编上 "1"
            self.assertFalse(h1s[0].paragraph_format.page_break_before)
            self.assertEqual("1\t问题重述", h1s[1].text)
            self.assertTrue(h1s[1].paragraph_format.page_break_before)

    def test_same_second_second_export_does_not_overwrite(self):
        """同秒重复导出: 冻结时钟跑两次, 第二次换名 _2 且旧产物字节不变 (v3.1.0)。"""
        import contextlib
        import io
        from datetime import datetime as _dt
        from unittest.mock import patch

        class _FrozenDatetime(_dt):
            @classmethod
            def now(cls, tz=None):
                return cls(2026, 9, 19, 12, 0, 0)

        with tempfile.TemporaryDirectory() as tmp:
            ws = Path(tmp) / "paper_workspace"
            out_dir = Path(tmp) / "submission"
            _write(ws / "01_abstract.md", "摘要正文。\n")
            with patch.object(export_final_docx, "datetime", _FrozenDatetime):
                buf = io.StringIO()
                with contextlib.redirect_stdout(buf):
                    rc1 = main(["--workspace", str(ws), "--out-dir", str(out_dir),
                                "--title", "T"])
                    self.assertEqual(0, rc1, buf.getvalue())
                    first = out_dir / "paper_final_20260919_120000.docx"
                    self.assertTrue(first.is_file())
                    first_bytes = first.read_bytes()
                    buf2 = io.StringIO()
                    with contextlib.redirect_stdout(buf2):
                        rc2 = main(["--workspace", str(ws), "--out-dir", str(out_dir),
                                    "--title", "T"])
                    self.assertEqual(0, rc2, buf2.getvalue())
                    self.assertIn("已存在 (同秒重复导出或并发导出)", buf2.getvalue())
            second = out_dir / "paper_final_20260919_120000_2.docx"
            self.assertTrue(second.is_file())
            self.assertEqual(first_bytes, first.read_bytes())  # 先到的产物未被覆盖
            self.assertEqual(2, len(list(out_dir.glob("*_final_*.docx"))))


if __name__ == "__main__":
    unittest.main()

class FakePandocPlaceholderTest(unittest.TestCase):
    """P2: pandoc 假成功 (rc0 无输出 / 0 字节) 必须报错且不留空终稿占位。

    注意 (事故教训): 假 run 必须**只**按 `-o 目标` 落盘, 不能盲写 cmd[-1] ——
    main 里的 `has_pandoc()` 会先探 `pandoc --version`, 盲写会在当前目录留下
    名为 `--version` 的空文件 (2026-09-19 实测, 已迁出保全)。
    """

    @staticmethod
    def _out_target(cmd):
        """从命令行取 `-o` 之后的目标; 找不到返回 None (探测类命令一律不落盘)。"""
        if "-o" not in cmd:
            return None
        return Path(cmd[cmd.index("-o") + 1])

    def _run(self, tmp: str, fake):
        ws = Path(tmp) / "paper_workspace"
        out_dir = Path(tmp) / "submission"
        _write(ws / "01_abstract.md", "摘要正文。\n")
        original = export_final_docx.subprocess.run
        export_final_docx.subprocess.run = fake
        try:
            import contextlib, io
            buf = io.StringIO()
            with contextlib.redirect_stdout(buf):
                rc = main(["--workspace", str(ws), "--out-dir", str(out_dir),
                           "--title", "T"])
        finally:
            export_final_docx.subprocess.run = original
        return rc, buf.getvalue(), out_dir

    def test_rc0_without_output_cleans_placeholder(self):
        from types import SimpleNamespace

        def _fake(cmd, **kwargs):
            return SimpleNamespace(returncode=0, stderr="", stdout="")

        with tempfile.TemporaryDirectory() as tmp:
            rc, out, out_dir = self._run(tmp, _fake)
            self.assertEqual(2, rc)
            self.assertIn("未生成有效 docx", out)
            self.assertEqual([], list(out_dir.glob("*_final_*.docx")))  # 无空占位残留

    def test_rc0_with_zero_byte_output_cleans_placeholder(self):
        from types import SimpleNamespace

        def _fake(cmd, **kwargs):
            target = self._out_target(cmd)
            if target is not None:              # 只写 -o 目标 (探测命令不落盘)
                target.write_bytes(b"")
            return SimpleNamespace(returncode=0, stderr="", stdout="")

        with tempfile.TemporaryDirectory() as tmp:
            rc, out, out_dir = self._run(tmp, _fake)
            self.assertEqual(2, rc)
            self.assertEqual([], list(out_dir.glob("*_final_*.docx")))


if __name__ == "__main__":
    unittest.main()
