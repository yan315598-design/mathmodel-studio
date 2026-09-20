"""Inspect local attachments with lazy optional readers; never modify inputs."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
import subprocess
import wave
from pathlib import Path


KINDS = {
    "table": {".csv", ".tsv", ".xlsx", ".xls", ".parquet"},
    "image": {".png", ".jpg", ".jpeg", ".bmp", ".webp"},
    "raster": {".tif", ".tiff", ".geotiff"},
    "video": {".mp4", ".avi", ".mov", ".mkv", ".webm"},
    "audio": {".wav", ".flac", ".mp3", ".ogg", ".m4a"},
    "vector": {".geojson", ".shp", ".gpkg"},
    "point_cloud": {".las", ".laz", ".ply", ".pcd"},
    "scientific": {".nc", ".h5", ".hdf5", ".npy", ".npz", ".mat"},
    "document": {".pdf", ".docx", ".txt", ".md"},
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def media_probe(path: Path) -> dict:
    if not shutil.which("ffprobe"):
        raise ImportError("ffprobe (FFmpeg)")
    result = subprocess.run(
        ["ffprobe", "-v", "error", "-show_format", "-show_streams", "-of", "json", str(path)],
        capture_output=True, text=True, encoding="utf-8", timeout=30, check=True,
    )
    raw = json.loads(result.stdout)
    fields = ("index", "codec_type", "codec_name", "width", "height", "sample_rate",
              "channels", "duration", "avg_frame_rate", "r_frame_rate", "nb_frames", "time_base")
    return {"duration_s": raw.get("format", {}).get("duration"),
            "streams": [{k: stream[k] for k in fields if k in stream} for stream in raw.get("streams", [])]}


def metadata(path: Path, kind: str) -> dict:
    suffix = path.suffix.lower()
    if suffix == ".wav":
        with wave.open(str(path), "rb") as audio:
            return {"sample_rate_hz": audio.getframerate(), "channels": audio.getnchannels(),
                    "frames": audio.getnframes(), "sample_width_bytes": audio.getsampwidth(),
                    "duration_s": audio.getnframes() / audio.getframerate()}
    if kind in {"audio", "video"}:
        return media_probe(path)
    if kind == "image":
        from PIL import Image
        with Image.open(path) as image:
            image.verify()
        with Image.open(path) as image:
            return {"width": image.width, "height": image.height, "mode": image.mode,
                    "format": image.format}
    if kind == "raster":
        with path.open("rb") as stream:
            if stream.read(4)[:2] not in {b"II", b"MM"}:
                raise ValueError("invalid TIFF byte order/header")
        georeferenced = suffix == ".geotiff"
        if not georeferenced:
            from PIL import Image, UnidentifiedImageError
            try:
                with Image.open(path) as image:
                    tags = image.tag_v2
                    georeferenced = any(tag in tags for tag in (33550, 33922, 34264, 34735))
                    if not georeferenced:
                        image.verify()
            except UnidentifiedImageError:
                georeferenced = True  # Scientific TIFF variants may require GDAL.
        if not georeferenced:
            with Image.open(path) as image:
                return {"width": image.width, "height": image.height, "mode": image.mode,
                        "format": image.format, "classification": "ordinary_tiff",
                        "geospatial_status": "no embedded georeferencing; sidecars not interpreted",
                        "reader": "Pillow"}
        import rasterio
        with rasterio.open(path) as dataset:
            return {"width": dataset.width, "height": dataset.height, "bands": dataset.count,
                    "crs": str(dataset.crs) if dataset.crs else None,
                    "bounds": list(dataset.bounds), "transform": list(dataset.transform),
                    "dtypes": list(dataset.dtypes), "nodata": str(dataset.nodata),
                    "units": list(dataset.units), "tags": dataset.tags(),
                    "reader": "rasterio", "classification": "geospatial_raster"}
    if kind == "table" and suffix in {".csv", ".tsv"}:
        with path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.reader(stream, delimiter="\t" if suffix == ".tsv" else ",")
            columns = next(reader)
            sample = []
            for _, row in zip(range(100), reader):
                sample.append(len(row))
        return {"columns": columns, "sample_rows": len(sample),
                "sample_width_consistent": all(n == len(columns) for n in sample),
                "scope": "header and first 100 rows; not a full data audit"}
    if suffix == ".xlsx":
        import openpyxl
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return {"sheets": [{"name": sheet.title, "rows": sheet.max_row,
                                "columns": sheet.max_column} for sheet in book]}
        finally:
            book.close()
    if suffix == ".parquet":
        import pyarrow.parquet as pq
        info = pq.read_metadata(path)
        return {"rows": info.num_rows, "columns": info.num_columns, "schema": str(info.schema)}
    if suffix in {".las", ".laz"}:
        import laspy
        with laspy.open(path) as cloud:
            return {"points": cloud.header.point_count, "mins": cloud.header.mins.tolist(),
                    "maxs": cloud.header.maxs.tolist(), "scales": cloud.header.scales.tolist(),
                    "offsets": cloud.header.offsets.tolist(), "scope": "header only"}
    if kind == "vector":
        import fiona
        with fiona.open(path) as dataset:
            return {"features": len(dataset), "crs": str(dataset.crs),
                    "bounds": list(dataset.bounds), "schema": dict(dataset.schema)}
    if suffix == ".npy":
        import numpy as np
        array = np.load(path, mmap_mode="r", allow_pickle=False)
        return {"shape": list(array.shape), "dtype": str(array.dtype)}
    if suffix == ".npz":
        import numpy as np
        with np.load(path, allow_pickle=False) as store:
            arrays = {}
            for name in store.files[:100]:
                array = store[name]
                arrays[name] = {"shape": list(array.shape), "dtype": str(array.dtype)}
            return {"arrays": arrays,
                    "scope": "first 100 arrays; compressed arrays decompressed", "reader": "numpy"}
    if suffix in {".h5", ".hdf5"}:
        import h5py
        items = []
        with h5py.File(path, "r") as dataset:
            def visit(name, value):
                if isinstance(value, h5py.Dataset):
                    items.append({"name": name, "shape": list(value.shape), "dtype": str(value.dtype)})
                return True if len(items) >= 100 else None
            dataset.visititems(visit)
        return {"datasets": items, "scope": "at most 100 dataset headers"}
    if suffix == ".nc":
        import xarray as xr
        with xr.open_dataset(path, decode_times=False) as dataset:
            return {"reader": "xarray", "attributes": {str(k): str(v) for k, v in dataset.attrs.items()},
                    "dimensions": dict(dataset.sizes), "variables": {
                name: {"dims": list(value.dims), "shape": list(value.shape), "dtype": str(value.dtype),
                       "units": str(value.attrs.get("units", "unknown")),
                       "attributes": {str(k): str(v) for k, v in value.attrs.items()}}
                for name, value in list(dataset.variables.items())[:100]}}
    if suffix == ".mat":
        from scipy.io import whosmat
        return {"variables": whosmat(path)}
    raise NotImplementedError("recognized type; use the task-specific reader")


def inspect(path: Path) -> dict:
    kind = next((name for name, extensions in KINDS.items() if path.suffix.lower() in extensions), "unknown")
    record = {"path": str(path.resolve()), "kind": kind, "status": "recognized",
              "source": "local_attachment", "metadata": {}}
    try:
        record.update(bytes=path.stat().st_size, sha256=sha256(path))
        record["asset_id"] = "asset_" + record["sha256"][:16]
        record["metadata"] = metadata(path, kind)
        record["status"] = "inspected"
    except ImportError as error:
        record.update(status="missing_dependency", detail=str(error))
    except NotImplementedError as error:
        record.update(status="unsupported_reader", detail=str(error))
    except Exception as error:
        record.update(status="error", detail=f"{type(error).__name__}: {error}")
    record["capability"] = {"metadata": record["status"], "analysis": "not_run"}
    record["provenance"] = {"input_sha256": record.get("sha256"), "derived_from": [],
                            "scope": "original local attachment; no transformation"}
    return record


def extract_frame(path: Path, index: int, output: Path, scan_limit: int = 10000) -> dict:
    """Use decoded frame order and original presentation timestamps, including VFR."""
    if not 0 <= index < scan_limit:
        raise ValueError("frame index must be within the explicit scan limit")
    if output.exists():
        raise ValueError("output already exists")
    if not all(shutil.which(tool) for tool in ("ffmpeg", "ffprobe")):
        raise ImportError("ffmpeg and ffprobe are required")
    probe = subprocess.run(
        ["ffprobe", "-v", "error", "-select_streams", "v:0", "-read_intervals", f"%+#{scan_limit}",
         "-show_frames", "-show_entries", "frame=best_effort_timestamp_time", "-of", "json", str(path)],
        capture_output=True, text=True, encoding="utf-8", timeout=60, check=True,
    )
    frames = json.loads(probe.stdout)["frames"]
    if index >= len(frames) or "best_effort_timestamp_time" not in frames[index]:
        raise ValueError("requested decoded frame or presentation timestamp unavailable")
    output.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["ffmpeg", "-v", "error", "-n", "-i", str(path), "-map", "0:v:0",
                    "-vf", f"select=eq(n\\,{index})", "-frames:v", "1", str(output)],
                   capture_output=True, timeout=60, check=True)
    if not output.is_file():
        raise ValueError("frame extraction produced no output")
    return {"source": str(path.resolve()), "source_sha256": sha256(path), "frame_index_zero_based": index,
            "timestamp_s": float(frames[index]["best_effort_timestamp_time"]), "scan_limit": scan_limit,
            "output": str(output.resolve()), "output_sha256": sha256(output), "backend": "ffmpeg"}


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("paths", nargs="+", type=Path)
    parser.add_argument("--output", type=Path, required=True, help="new JSON manifest")
    parser.add_argument("--frame", type=int, help="extract one zero-based video frame")
    parser.add_argument("--image", type=Path)
    parser.add_argument("--scan-limit", type=int, default=10000)
    args = parser.parse_args()
    if args.output.exists():
        parser.error("output exists; use a new manifest path")
    if args.frame is not None:
        if len(args.paths) != 1 or args.image is None:
            parser.error("frame extraction requires one video and --image")
        if args.image.resolve() == args.output.resolve():
            parser.error("image and JSON manifest must use different paths")
        try:
            data = {"schema_version": "asset-frame-1", "frame": extract_frame(
                args.paths[0], args.frame, args.image, args.scan_limit)}
        except (ValueError, ImportError, OSError, subprocess.SubprocessError) as error:
            parser.exit(2, f"[ERROR] {error}\n")
        failures = 0
    else:
        records = [inspect(path) for path in args.paths]
        data = {"schema_version": "assets-1", "assets": records}
        failures = sum(item["status"] != "inspected" for item in records)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("x", encoding="utf-8") as stream:
        json.dump(data, stream, ensure_ascii=False, indent=2, allow_nan=False)
    print(f"[assets] output={args.output} incomplete={failures}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
